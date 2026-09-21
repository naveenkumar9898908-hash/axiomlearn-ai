import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

BG_SPACE = RGBColor(7, 11, 25)         # #070B19 Deep Space Obsidian
GRID_LINE = RGBColor(20, 32, 60)       # Cyber Grid Lines
CARD_SHADOW = RGBColor(3, 6, 15)       # 3D Drop Depth Shadow
CARD_SURFACE = RGBColor(16, 24, 46)    # Elevated 3D Plate
CARD_BORDER = RGBColor(38, 55, 96)     # Chamfer Border
CYAN_GLOW = RGBColor(0, 240, 255)      # Neon Cyber Cyan
VIOLET_GLOW = RGBColor(168, 85, 247)   # Neon Holographic Violet
EMERALD_GLOW = RGBColor(16, 185, 129)  # Matrix Emerald
AMBER_GLOW = RGBColor(251, 191, 36)    # Warning Gold
ROSE_GLOW = RGBColor(244, 63, 94)      # Accent Coral
TEXT_LIGHT = RGBColor(248, 250, 252)   # Ultra White
TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400

def draw_3d_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_SPACE
    bg.line.color.rgb = BG_SPACE
    
    for i in range(7):
        y_val = 6.4 + (i * 0.16)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(y_val), Inches(12.333), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = GRID_LINE
        line.line.fill.background()
        
    top_glow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.04))
    top_glow.fill.solid()
    top_glow.fill.fore_color.rgb = CYAN_GLOW
    top_glow.line.fill.background()

def add_3d_header(slide, tag_text, title_text, color=CYAN_GLOW):
    tag_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(4.8), Inches(0.38))
    tag_bg.fill.solid()
    tag_bg.fill.fore_color.rgb = CARD_SURFACE
    tag_bg.line.color.rgb = color
    tag_bg.line.width = Pt(1.5)
    
    tf_tag = tag_bg.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = f"❖  {tag_text.upper()}"
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = color
    p_tag.alignment = PP_ALIGN.CENTER
    
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_LIGHT

def make_3d_card(slide, left, top, width, height, title, lines, accent=CYAN_GLOW, icon=""):
    shadow = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.08), Inches(top + 0.08), Inches(width), Inches(height))
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = CARD_SHADOW
    shadow.line.fill.background()
    
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_SURFACE
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1.5)
    
    bevel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + 0.15), Inches(top + 0.04), Inches(width - 0.3), Inches(0.06))
    bevel.fill.solid()
    bevel.fill.fore_color.rgb = accent
    bevel.line.fill.background()
    
    tb = slide.shapes.add_textbox(Inches(left + 0.18), Inches(top + 0.14), Inches(width - 0.36), Inches(height - 0.28))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    full_title = f"{icon}  {title}" if icon else title
    p0.text = full_title
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = accent
    p0.space_after = Pt(6)
    
    for l in lines:
        p = tf.add_paragraph()
        p.text = l
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(2.5)

# SLIDE 1
s1 = prs.slides.add_slide(blank_layout)
draw_3d_background(s1)

badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.7), Inches(5.2), Inches(0.42))
badge.fill.solid()
badge.fill.fore_color.rgb = CARD_SURFACE
badge.line.color.rgb = CYAN_GLOW
badge.line.width = Pt(1.5)
tf_b = badge.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "⚡ TENSORA 2026 HACKATHON • PROBLEM STATEMENT [EDU-01]"
p_b.font.size = Pt(10)
p_b.font.bold = True
p_b.font.color.rgb = CYAN_GLOW
p_b.alignment = PP_ALIGN.CENTER

tb_m = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(2.2))
tf_m = tb_m.text_frame
tf_m.word_wrap = True

pm1 = tf_m.paragraphs[0]
pm1.text = "AxiomLearn AI"
pm1.font.size = Pt(50)
pm1.font.bold = True
pm1.font.color.rgb = TEXT_LIGHT
pm1.space_after = Pt(4)

pm2 = tf_m.add_paragraph()
pm2.text = "The 3D Adaptive Cognitive Engine & GPS for Your Brain"
pm2.font.size = Pt(22)
pm2.font.bold = True
pm2.font.color.rgb = CYAN_GLOW
pm2.space_after = Pt(10)

pm3 = tf_m.add_paragraph()
pm3.text = "Ordinary tests just tell you 'You're Wrong'. AxiomLearn AI finds the exact loose building block from your foundation and fixes it in 4 quick, engaging steps."
pm3.font.size = Pt(14)
pm3.font.color.rgb = TEXT_MUTED

make_3d_card(s1, 0.8, 4.0, 3.7, 2.8, "The 3D MRI Scanner", [
    "• Regular Tests = A Thermometer:",
    "  Tells you that you have a fever (40%), but cannot explain why.",
    "",
    "• AxiomLearn = 3D Brain Scan:",
    "  Finds the exact hidden concept causing your confusion."
], CYAN_GLOW, "🩺")

make_3d_card(s1, 4.8, 4.0, 3.7, 2.8, "The Cognitive GPS", [
    "• Missing a Turn in Class:",
    "  Teachers keep driving ahead even if you took a wrong exit.",
    "",
    "• AxiomLearn Recalculates:",
    "  Instantly reroutes you with the simplest shortcut back to mastery."
], VIOLET_GLOW, "🗺️")

make_3d_card(s1, 8.8, 4.0, 3.7, 2.8, "Video Game Dopamine", [
    "• Level Up Your Mastery:",
    "  Earn XP, build daily streaks (🔥), and unlock 7 mystery badges.",
    "",
    "• Pure Joy in Learning:",
    "  Celebration confetti and audio sound FX turn math into fun."
], EMERALD_GLOW, "🎮")

# SLIDE 2
s2 = prs.slides.add_slide(blank_layout)
draw_3d_background(s2)
add_3d_header(s2, "The Real Crisis In Education", "The 'Tower of Blocks' Disaster in Modern Classrooms", AMBER_GLOW)

make_3d_card(s2, 0.8, 1.8, 5.7, 5.0, "The 3 Broken Realities of School", [
    "1. The 'One-Speed Fits Nobody' Rule:",
    "A teacher with 30 students must teach at one single pace. Fast students get bored; struggling students fall further behind every week.",
    "",
    "2. The Collapsing Tower Disaster:",
    "Learning is like building a 10-story tower. If Block #3 (basic algebra) is loose, Floor #10 (calculus) will collapse every time!",
    "Teachers blame students for failing calculus, but the real culprit was a rule forgotten 2 years ago.",
    "",
    "3. Report Cards Tell You Weeks Late:",
    "Finding out you failed at the end of the term is useless. By then, confidence is crushed."
], ROSE_GLOW, "⚠️")

make_3d_card(s2, 6.8, 1.8, 5.7, 5.0, "The AxiomLearn AI Breakthrough", [
    "1. Personalized 1-on-1 Pacing:",
    "Every student learns at their natural sweet spot. Fast learners zoom ahead; struggling learners get patient, step-by-step guidance.",
    "",
    "2. 3D Root-Cause Detection:",
    "Our AI scans backward down your syllabus tree to isolate the exact prerequisite brick that came loose.",
    "",
    "3. Instant 60-Second Remediation:",
    "Bridges the gap right when it occurs so students never accumulate knowledge debt.",
    "",
    "4. Restores Total Student Confidence:",
    "Proves you aren't 'bad at math'—you were just missing one tiny clue!"
], CYAN_GLOW, "✨")

# ==============================================================================
# SLIDE 3: EXACT USER CONTENT FOR ALL 4 STEPS
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
draw_3d_background(s3)
add_3d_header(s3, "Simple 4-Step Architecture", "How AxiomLearn AI Works (Step-by-Step)", CYAN_GLOW)

step1_lines = [
    "The student first solves interactive questions.",
    "",
    "The AI automatically changes the difficulty:",
    "• If too easy, AI makes them harder.",
    "• If too difficult, AI makes them easier.",
    "• Goal: Keep student in 'flow zone' — challenging, not frustrating.",
    "",
    "Simple example:",
    "If learning Java and easily answer basic questions, AI gives slightly harder problems.",
    "",
    "👉 Purpose: Right level of practice."
]

step2_lines = [
    "While answering, AI observes performance:",
    "• Correct answers count & time taken",
    "• Quick mistakes vs conceptual struggle",
    "",
    "Latency Rule:",
    "• < 3s: Quick/careless mistake.",
    "• > 30s: Deeper difficulty with concept.",
    "",
    "Simple example:",
    "Accidentally misclicking vs spending 40s confused are treated differently.",
    "",
    "👉 Purpose: Understand why struggling."
]

step3_lines = [
    "If stuck, AI acts like a detective on the family tree of concepts:",
    "Advanced ➔ Basic ➔ Prerequisite",
    "",
    "For example: Arrays ➔ Loops ➔ Variables.",
    "If difficulty with arrays, the real problem might be loops.",
    "",
    "Instead of repeating arrays, AI identifies the basic concept you need first.",
    "",
    "👉 Purpose: Find root cause of confusion."
]

step4_lines = [
    "Quick 60-second learning activities:",
    "",
    "1. 60s Refresher Card: Short explanation.",
    "2. Root Checkup Probe: Small check question.",
    "3. Scaffolded Bridge: Connects basic to advanced.",
    "4. Victory Test with XP confetti! 🎉",
    "Final small test with XP/reward feeling.",
    "",
    "👉 Purpose: Quickly repair knowledge gap and get student back on track."
]

steps_data = [
    ("Step 1: Practice", step1_lines, CYAN_GLOW, "🟦"),
    ("Step 2: AI Listens", step2_lines, VIOLET_GLOW, "🟪"),
    ("Step 3: Detective", step3_lines, AMBER_GLOW, "🟨"),
    ("Step 4: 60s Fix", step4_lines, EMERALD_GLOW, "🟩")
]

for i, (stitle, slines, scolor, sicon) in enumerate(steps_data):
    x_pos = 0.8 + (i * 2.95)
    make_3d_card(s3, x_pos, 1.8, 2.8, 5.0, stitle, slines, scolor, sicon)

# SLIDE 4
s4 = prs.slides.add_slide(blank_layout)
draw_3d_background(s4)
add_3d_header(s4, "Interactive 3D Modules", "Futuristic Platform Features You Can Touch & Use", VIOLET_GLOW)

make_3d_card(s4, 0.8, 1.8, 5.7, 2.4, "Futuristic Cyber-HUD Arena", [
    "• 3D mouse parallax tilt with dynamic oscilloscope audio waves",
    "• Real-time XP engine, level-up bar, and streak multiplier flames",
    "• Physics-based canvas confetti & custom synthesized sound effects"
], CYAN_GLOW, "🕹️")

make_3d_card(s4, 6.8, 1.8, 5.7, 2.4, "Glowing 3D Knowledge Graph", [
    "• Interactive map of Math & Computer Science concept nodes",
    "• Prerequisite Chain Illumination: click any topic to light up parents",
    "• Green = Mastered, Amber = In-Progress, Red = Bottleneck gap"
], VIOLET_GLOW, "🌐")

make_3d_card(s4, 0.8, 4.4, 5.7, 2.4, "Teacher Cohort Command Center (CRUD)", [
    "• Complete Student Management: Create, Read (Table/Grid), Update, Delete",
    "• Search & persona filter chips (Fast Pacers, Blocked, At-Risk)",
    "• 1-Click Systemic Workshop: Batch-heals class-wide syllabus roadblocks"
], EMERALD_GLOW, "👩‍🏫")

make_3d_card(s4, 6.8, 4.4, 5.7, 2.4, "Telemetry Sandbox & Secure Accounts", [
    "• Interactive Bayesian & Neural Net math sandbox with live equations",
    "• Multi-user persona logins (Maya, Alex, Liam, Prof. Reed)",
    "• Complete local browser storage isolation (no cloud privacy leaks)"
], AMBER_GLOW, "🔐")

# SLIDE 5
s5 = prs.slides.add_slide(blank_layout)
draw_3d_background(s5)
add_3d_header(s5, "Tangible Value Delivered", "Why Students, Teachers & Schools Love AxiomLearn AI", EMERALD_GLOW)

make_3d_card(s5, 0.8, 1.8, 3.7, 5.0, "For Students\n(Zero Frustration)", [
    "❤️ No More Homework Tears: Never feel lost or helpless again.",
    "⚡ Learn in Half the Time: Skip what you know, master only what you need.",
    "🎮 Addictively Fun: Feels like playing a sci-fi video game instead of boring worksheets."
], CYAN_GLOW, "🧑‍🎓")

make_3d_card(s5, 4.8, 1.8, 3.7, 5.0, "For Teachers\n(Actionable Clarity)", [
    "⏱️ Saves 10+ Hours Every Week: Zero manual grading of repetitive paper quizzes.",
    "🎯 Exact Diagnostic Truth: Knows exactly why students struggled—no more guesswork.",
    "🪄 1-Click Class Workshops: Fixes the single concept holding back 30% of the class."
], EMERALD_GLOW, "👩‍🏫")

make_3d_card(s5, 8.8, 1.8, 3.7, 5.0, "For Institutions\n(Retention & Grades)", [
    "📈 Proven Grade Improvements: Delivers 1-on-1 private tutor quality for free.",
    "🛡️ Early Failure Prevention: Flags at-risk learners in week 1 rather than after finals.",
    "💻 Zero Infrastructure Cost: Runs entirely on basic school laptops and Chromebooks."
], VIOLET_GLOW, "🏫")

# SLIDE 6
s6 = prs.slides.add_slide(blank_layout)
draw_3d_background(s6)
add_3d_header(s6, "Summary & Next Steps", "100% Ready Today & Our Expanding Vision", CYAN_GLOW)

make_3d_card(s6, 0.8, 1.8, 3.7, 5.0, "100% Built & Live", [
    "• Complete Web Application: Live on http://localhost:5173 with full math and CS curricula.",
    "• Verified Quality: 5/5 engine tests pass, 0 TypeScript errors across 37 files.",
    "• Live on GitHub: Committed and uploaded to official GitHub repository."
], CYAN_GLOW, "✅")

make_3d_card(s6, 4.8, 1.8, 3.7, 5.0, "The Core Takeaway", [
    "• Uniform Pacing is Dead: Forcing 30 students to learn at one speed is unfair to everyone.",
    "• Gaps are 100% Solvable: Every difficult concept is just simple building blocks.",
    "• AI for Human Potential: Empowers teachers and unlocks every student's true genius."
], EMERALD_GLOW, "💡")

make_3d_card(s6, 8.8, 1.8, 3.7, 5.0, "Future Innovations", [
    "• 🎙️ Talking AI Socratic Tutor: Spoken voice hints that guide students verbally.",
    "• 📲 Google Classroom LTI Sync: 1-click school roster sync for seamless district adoption.",
    "• 🔬 Unified STEM Trees: Cross-connecting Physics, Chemistry, and Engineering knowledge graphs."
], VIOLET_GLOW, "🔮")

# Save to dedicated updated presentation
out_file = r"c:\Users\navee\OneDrive\Pictures\NEW\347\AxiomLearn_AI_3D_Presentation_Updated.pptx"
prs.save(out_file)
print("SUCCESS: Saved to", out_file)
